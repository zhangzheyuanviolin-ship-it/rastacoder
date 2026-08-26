#!/usr/bin/env python3
import os
import sys
import tempfile
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'python'))


def require(condition, message):
    if not condition:
        raise AssertionError(message)


# ---------------------------------------------------------------------------
# Source/UI invariants
# ---------------------------------------------------------------------------
skill_text = (ROOT / 'lib/core/models/tool_skill.dart').read_text()
require("RASTACODER_V7_COMPLETE_SKILLS" in skill_text, 'v7 skill marker missing')
require("upstreamExtendedToolNames" in skill_text, 'upstream post-baseline tool set missing')
require("v7AddedToolNames" in skill_text, 'v7 added tool set missing')
for capability in (
    '删除文件/目录', '安全解压', '页面重排', '尺寸放大/缩小',
    '视频拼接', '多音轨混音', '高级 custom FFmpeg', '新建 PPTX', '新建 XLSX',
):
    require(capability in skill_text, f'skill capability missing: {capability}')

skills_ui = (ROOT / 'lib/features/settings/tool_skills_screen.dart').read_text()
require('规范工具覆盖' in skills_ui, 'stale 23-tool wording remains')
require("skill.capabilities.join('、')" in skills_ui, 'screen reader does not announce skill actions')

chat = (ROOT / 'lib/features/chat/presentation/chat_screen.dart').read_text()
require("label: '聊天记录'" in chat and "label: '新建对话'" in chat, 'chat app-bar semantic labels missing')

bubble = (ROOT / 'lib/features/chat/presentation/widgets/message_bubble.dart').read_text()
require('explicitChildNodes: true' in bubble, 'nested Thinking/diagnostics semantics not preserved')
require('excludeSemantics: true' not in bubble, 'message bubble still swallows interactive semantics')
require('思考过程，当前已折叠，双击展开' in bubble, 'Thinking control label missing')
require('工具调用诊断，当前已折叠，双击展开' in bubble, 'diagnostics control label missing')

storage = (ROOT / 'lib/core/services/storage_service.dart').read_text()
require("return value ?? 'qwen3-4b';" in storage, 'fresh-install model default is not local Qwen3')
main = (ROOT / 'lib/main.dart').read_text()
require('RASTACODER_V7_LOCAL_MODEL_RESTORE' in main, 'cold-start local model restore missing')
require('loadModel(preferredModel)' in main, 'preferred local model not loaded at startup')

native = (ROOT / 'lib/core/services/native_tool_executor.dart').read_text()
for op in ("case 'concat':", "case 'mix_audio':", "case 'merge_av':", "case 'custom':"):
    require(op in native, f'FFmpeg operation missing: {op}')
require("args['input_paths']" in native, 'native multi-input FFmpeg paths missing')


# ---------------------------------------------------------------------------
# Tool registry completeness: 21 manual Skills may overlap, but their union must
# exactly cover the model-facing local canonical schema. No stale "23" shortcut.
# ---------------------------------------------------------------------------
from navixmind.tools import (  # noqa: E402
    OFFLINE_TOOLS_SCHEMA,
    LOCAL_SKILLS,
    get_enabled_tool_names,
    get_offline_tools_for_skills,
)

schema_names = {item['name'] for item in OFFLINE_TOOLS_SCHEMA}
covered = get_enabled_tool_names(tuple(LOCAL_SKILLS.keys()))
require(len(LOCAL_SKILLS) == 21, f'expected 21 manual Skills, got {len(LOCAL_SKILLS)}')
require(schema_names == covered, f'local schema/Skill coverage mismatch: schema-only={schema_names-covered}, skill-only={covered-schema_names}')
require(len(schema_names) == 31, f'expected 31 canonical local tools after v7 expansion, got {len(schema_names)}')
for name in (
    'list_files', 'file_manage', 'list_zip', 'extract_zip', 'pdf_manage',
    'image_compose', 'create_pptx', 'create_xlsx', 'python_execute', 'ffmpeg_process',
):
    require(name in schema_names, f'canonical local tool missing: {name}')

ffmpeg = next(item for item in OFFLINE_TOOLS_SCHEMA if item['name'] == 'ffmpeg_process')
ops = set(ffmpeg['input_schema']['properties']['operation']['enum'])
for op in ('custom', 'concat', 'mix_audio', 'merge_av'):
    require(op in ops, f'local FFmpeg schema hides operation: {op}')
require('input_paths' in ffmpeg['input_schema']['properties'], 'local FFmpeg multi-input schema missing')

text_skill = {x['name'] for x in get_offline_tools_for_skills(['text_files'])}
require({'read_file', 'write_file', 'file_info', 'list_files', 'file_manage'} <= text_skill, 'text/file Skill is still incomplete')
zip_skill = {x['name'] for x in get_offline_tools_for_skills(['zip_archive'])}
require({'create_zip', 'list_zip', 'extract_zip'} <= zip_skill, 'ZIP Skill lacks create/list/extract')
image_skill = {x['name'] for x in get_offline_tools_for_skills(['image_processing'])}
require({'image_compose', 'smart_crop'} <= image_skill, 'image Skill lacks full image processing surface')


# ---------------------------------------------------------------------------
# Functional smoke tests for new structured tools.
# ---------------------------------------------------------------------------
from navixmind.tools.extended_tools import (  # noqa: E402
    list_files, file_manage, list_zip, extract_zip, pdf_manage,
    create_pptx, create_xlsx, image_compose,
)
from PIL import Image  # noqa: E402
from pypdf import PdfWriter, PdfReader  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from pptx import Presentation  # noqa: E402

with tempfile.TemporaryDirectory() as td:
    root = Path(td)
    output = root / 'output'
    output.mkdir()

    # file lifecycle: mkdir -> touch -> list -> copy -> move -> delete
    folder = file_manage('mkdir', path='folder', _output_dir=str(output))['path']
    touched = file_manage('touch', path='folder/a.txt', _output_dir=str(output))['path']
    Path(touched).write_text('hello')
    listing = list_files(path=folder)
    require(any(e['name'] == 'a.txt' for e in listing['entries']), 'list_files did not discover created file')
    copied = file_manage('copy', source_path=touched, destination_path='folder/b.txt', _output_dir=str(output))['destination_path']
    moved = file_manage('move', source_path=copied, destination_path='folder/c.txt', _output_dir=str(output))['destination_path']
    require(Path(moved).exists(), 'move failed')
    file_manage('delete', path=moved)
    require(not Path(moved).exists(), 'delete failed')

    # ZIP list/extract.
    zip_path = output / 'sample.zip'
    with zipfile.ZipFile(zip_path, 'w') as zf:
        zf.writestr('inside.txt', 'zip works')
    require(list_zip(str(zip_path))['count'] == 1, 'list_zip failed')
    extracted = extract_zip(str(zip_path), _output_dir=str(output))
    require(Path(extracted['output_paths'][0]).read_text() == 'zip works', 'extract_zip failed')

    # Image resize + format output.
    img1 = output / 'one.png'
    img2 = output / 'two.png'
    Image.new('RGB', (100, 50), 'white').save(img1)
    Image.new('RGB', (80, 50), 'black').save(img2)
    resized = output / 'resized.jpg'
    image_compose([str(img1)], str(resized), 'resize', {'width': 400})
    with Image.open(resized) as im:
        require(im.width == 400 and im.height == 200, 'image resize did not preserve aspect ratio')
    joined = output / 'joined.png'
    image_compose([str(img1), str(img2)], str(joined), 'concat_horizontal', {})
    with Image.open(joined) as im:
        require(im.width == 180 and im.height == 50, 'image concat failed')

    # PDF split/merge/page operations.
    p1 = output / 'p1.pdf'
    p2 = output / 'p2.pdf'
    for path in (p1, p2):
        w = PdfWriter(); w.add_blank_page(width=100, height=100)
        with open(path, 'wb') as f: w.write(f)
    merged = pdf_manage('merge', input_paths=[str(p1), str(p2)], output_path='merged.pdf', _output_dir=str(output))
    require(len(PdfReader(merged['output_path']).pages) == 2, 'PDF merge failed')
    split = pdf_manage('split', input_path=merged['output_path'], _output_dir=str(output))
    require(len(split['output_paths']) == 2, 'PDF split failed')

    # PPTX/XLSX creation.
    ppt = output / 'deck.pptx'
    create_pptx(str(ppt), title='Title', slides=[{'title': 'Slide', 'content': 'Body'}])
    require(len(Presentation(str(ppt)).slides) == 2, 'create_pptx failed')
    xlsx = output / 'book.xlsx'
    create_xlsx(str(xlsx), sheets=[{'name': 'Data', 'rows': [[1, 2], [3, 4]]}])
    wb = load_workbook(xlsx)
    require(wb['Data']['B2'].value == 4, 'create_xlsx failed')
    wb.close()

print('RastaCoder v7 complete-skill validation passed: 21 Skills / 31 canonical tools / new functional smoke tests green.')
